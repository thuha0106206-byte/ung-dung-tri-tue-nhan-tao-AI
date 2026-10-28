import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

os.makedirs('docs/test_cases', exist_ok=True)

def style_excel(ws, title):
    # Header Title Banner
    ws.merge_cells('A1:J1')
    title_cell = ws['A1']
    title_cell.value = title
    title_cell.font = Font(name='Segoe UI', size=15, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40

    # Header Row
    headers = [
        'Mã TC', 'Phân Hệ / Module', 'Tên Ca Kiểm Thử', 'Điều Kiện Tiên Quyết',
        'Các Bước Thực Hiện', 'Dữ Liệu Đầu Vào (Input)', 'Kết Quả Mong Đợi (Expected)',
        'Mức Độ (Severity)', 'Kết Quả (Status)', 'Ghi Chú Kỹ Thuật'
    ]
    
    ws.append([]) # Row 2 empty
    ws.append(headers) # Row 3
    ws.row_dimensions[3].height = 28
    
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_num in range(1, 11):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

def format_data_rows(ws, data):
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )
    pass_fill = PatternFill(start_color='D1E7DD', end_color='D1E7DD', fill_type='solid')
    pass_font = Font(name='Segoe UI', size=10, bold=True, color='0F5132')
    
    crit_font = Font(name='Segoe UI', size=10, bold=True, color='842029')
    crit_fill = PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')

    for r_idx, row in enumerate(data, start=4):
        ws.append(row)
        ws.row_dimensions[r_idx].height = 42
        for c_idx in range(1, 11):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = Font(name='Segoe UI', size=10)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            
            # Center ID & Status & Severity
            if c_idx in [1, 8, 9]:
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
            if c_idx == 9 and cell.value == 'PASS':
                cell.fill = pass_fill
                cell.font = pass_font
                
            if c_idx == 8 and cell.value == 'Critical':
                cell.fill = crit_fill
                cell.font = crit_font

    # Set column widths
    widths = [12, 18, 28, 24, 32, 26, 32, 14, 12, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

# ----------------- WORKBOOK 1: KIỂM SOÁT AI & DUYỆT KTV -----------------
wb1 = openpyxl.Workbook()
ws1 = wb1.active
ws1.title = 'AI_KTV_Control_TestCases'
style_excel(ws1, 'BẢNG TEST CASE KIỂM THỬ: QUYỀN KIỂM SOÁT AI & PHÊ DUYỆT CỦA KỸ THUẬT VIÊN (FROGHIPPO PHONE)')

tc_ai = [
    [
        'TC_AI_01', 'AI Chẩn Đoán', 'Sinh bản thảo nội bộ (AI Draft Gating)',
        'Server đang chạy, KTV mở trang ky_thuat_vien.html',
        '1. Nhập Thiết bị & Triệu chứng lỗi\n2. Bấm "Phân Tích & Chẩn Đoán Lỗi"\n3. Kiểm tra hiển thị kết quả',
        'Dòng máy: iPhone 18 Pro Max\nLỗi: Rơi xe máy trượt 2m vỡ nát màn hình và sọc',
        'Hiện Banner Ổ Khóa Đỏ: 🔒 BẢN THẢO NỘI BỘ (AI DRAFT). Cờ daDuyetChuanDoan=False. Khách hàng và Chatbot bị chặn không thấy kết quả.',
        'Critical', 'PASS', 'Đã kiểm thử trên ky_thuat_vien.html'
    ],
    [
        'TC_AI_02', 'Kiểm Soát KTV', 'Hiệu chỉnh Nguyên nhân hư hỏng thực tế',
        'Đã có kết quả AI chẩn đoán sơ bộ (TC_AI_01)',
        '1. Đọc nội dung chẩn đoán của AI\n2. Click vào textarea "Nguyên Nhân Kỹ Thuật"\n3. Sửa/gõ lại nội dung theo máy đo kiểm',
        'Nhập: "Bể vỡ tấm nền OLED, đứt cổ cáp hiển thị do chấn động cơ học mạnh"',
        'KTV gõ và sửa trực tiếp được 100% câu chữ, không bị vô hiệu hóa hay khóa readonly.',
        'Critical', 'PASS', 'Đã tích hợp textarea cho phép sửa'
    ],
    [
        'TC_AI_03', 'Kiểm Soát KTV', 'Hiệu chỉnh Cảnh báo rủi ro tiềm ẩn',
        'Đã có kết quả AI chẩn đoán sơ bộ',
        '1. Click vào textarea "Cảnh Báo Rủi Ro Tiềm Ẩn"\n2. Chỉnh sửa nội dung cảnh báo thực tế',
        'Nhập: "Vết nứt kính để lâu làm bụi nước vào gây chết IC hiển thị và đứt mạch ngầm"',
        'KTV lưu giữ hoặc bổ sung rủi ro chuyên ngành chuẩn xác theo tình trạng máy.',
        'High', 'PASS', 'Textarea rủi ro hoạt động mượt'
    ],
    [
        'TC_AI_04', 'Kiểm Soát Giá', 'Hiệu chỉnh Đơn giá VNĐ & Thời hạn bảo hành',
        'AI đề xuất 2 gói sửa chữa',
        '1. Xem 2 gói sửa chữa gợi ý\n2. Sửa ô giá gói 1 và gói 2\n3. Sửa ô bảo hành',
        'Gói Tiết Kiệm: 1.250.000 đ (BH 6T)\nGói Chính Hãng: 2.800.000 đ (BH 12T)',
        'Hết lỗi 0 đ, ô input number nhận giá trị VNĐ chuẩn xác, cho phép KTV định giá linh kiện.',
        'Critical', 'PASS', 'Đã tạo ô input giá tiền riêng'
    ],
    [
        'TC_AI_05', 'Kiểm Soát KTV', 'Bác bỏ gợi ý của AI (Reject AI Draft)',
        'AI chẩn đoán sai hoặc phỏng đoán ảo giác (hallucination)',
        '1. Bấm nút [❌ Bác bỏ gợi ý AI]\n2. Kiểm tra giao diện và bộ nhớ tạm',
        'Bấm nút Bác bỏ',
        'Khối kết quả AI bị ẩn và xóa sạch hoàn toàn, hiển thị toast thông báo đã hủy bỏ gợi ý AI.',
        'High', 'PASS', 'Hàm handleRejectAIDiagnose hoạt động'
    ],
    [
        'TC_AI_06', 'Phê Duyệt KTV', 'Bấm nút Duyệt & Mở khóa công khai cho khách',
        'Đã chọn Phiếu sửa chữa và hoàn thiện nội dung chẩn đoán',
        '1. Chọn phiếu #PS01 từ dropdown\n2. Bấm [✅ KTV PHÊ DUYỆT & CÔNG KHAI CHO KHÁCH HÀNG]\n3. Kiểm tra DB & UI',
        'Mã phiếu: PS01, KTV: Lê Văn Cường (NV03)',
        'Lưu vào MySQL: daDuyetChuanDoan=True, nguoiDuyet="Lê Văn Cường". Banner chuyển sang xanh 🔓 ĐÃ MỞ KHÓA.',
        'Critical', 'PASS', 'API /api/ktv/duyet-chuan-doan/{ma_phieu} OK'
    ],
    [
        'TC_AI_07', 'Cổng Khách Hàng', 'Tra cứu phiếu ĐÃ DUYỆT chẩn đoán',
        'Phiếu #PS01 đã được KTV duyệt (TC_AI_06)',
        '1. Mở khach_hang.html\n2. Nhập mã "PS01" -> Bấm Tra Cứu\n3. Kiểm tra mục chẩn đoán',
        'Mã tra cứu: PS01',
        'Hiển thị khung xanh: ✅ [ĐÃ ĐƯỢC KTV LÊ VĂN CƯỜNG DUYỆT]: Hiển thị đúng nguyên nhân và phương án KTV đã duyệt.',
        'Critical', 'PASS', 'Đã kiểm tra API & UI khách hàng'
    ],
    [
        'TC_AI_08', 'Cổng Khách Hàng', 'Tra cứu phiếu CHƯA DUYỆT chẩn đoán',
        'Phiếu #PS02 đang ở trạng thái DangKiemTra, chưa được KTV duyệt',
        '1. Mở khach_hang.html\n2. Nhập mã "PS02" -> Bấm Tra Cứu\n3. Kiểm tra mục chẩn đoán',
        'Mã tra cứu: PS02',
        'Hiển thị khung vàng có ổ khóa: 🔒 [ĐANG CHỜ KTV THẨM ĐỊNH & DUYỆT]. Khách không thấy nội dung AI phỏng đoán.',
        'Critical', 'PASS', 'Đã test thành công trên PS02'
    ],
    [
        'TC_AI_09', 'AI Chatbot', 'Khách hỏi Chatbot về phiếu ĐÃ DUYỆT',
        'Phiếu #PS01 đã được duyệt',
        '1. Mở Chatbot tư vấn\n2. Hỏi: "Máy iPhone 13 Pro Max #PS01 của tôi bị sao vậy shop?"',
        'Câu hỏi: Máy #PS01 bị lỗi gì?',
        'Chatbot trả lời chuẩn xác dựa trên chẩn đoán KTV đã phê duyệt, cam kết không bịa đặt.',
        'High', 'PASS', 'Chatbot đọc field chuanDoanKTV'
    ],
    [
        'TC_AI_10', 'AI Chatbot', 'Khách hỏi Chatbot về phiếu CHƯA DUYỆT',
        'Phiếu #PS02 chưa được duyệt',
        '1. Mở Chatbot tư vấn\n2. Hỏi: "Máy #PS02 của tôi đã kiểm tra xong chưa, hỏng gì vậy?"',
        'Câu hỏi: Máy #PS02 hỏng gì?',
        'Chatbot từ chối tiết lộ nguyên nhân phỏng đoán, thông báo KTV đang đo kiểm trực tiếp và sẽ thông báo sau.',
        'Critical', 'PASS', 'Đã test logic fallback 3.1 & prompt'
    ]
]
format_data_rows(ws1, tc_ai)
file1_path = os.path.abspath('docs/test_cases/Bang_Test_Case_Kiem_Soat_AI_KTV.xlsx')
wb1.save(file1_path)
print(f"Generated: {file1_path}")


# ----------------- WORKBOOK 2: QUY CHUẨN SỬA CHỮA & CỌC 30/70 -----------------
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = 'Repair_Standard_Deposit_Cases'
style_excel(ws2, 'BẢNG TEST CASE KIỂM THỬ: QUY CHUẨN SỬA CHỮA, ĐẶT CỌC 30%/70% & BẢO HÀNH (FROGHIPPO PHONE)')

tc_core = [
    [
        'TC_STD_01', 'Tiếp Nhận Máy', 'Lập biên nhận tiếp nhận máy & Ghi nhận triệu chứng',
        'Khách mang máy đến quầy lễ tân',
        '1. Lễ tân nhập thông tin khách hàng & dòng máy\n2. Ghi nhận triệu chứng khách khai báo\n3. Xuất mã phiếu sửa',
        'Khách: Trần Văn B, SĐT: 0988776655, Máy: Samsung S22 Ultra, Lỗi: Chân sạc lỏng',
        'Hệ thống tạo phiếu trạng thái KhoiTao, in biên nhận có mã vạch/mã số tra cứu tức thì.',
        'High', 'PASS', 'Đã kiểm thử trên quầy tiếp nhận'
    ],
    [
        'TC_STD_02', 'Chính Sách Cọc', 'Thu cọc 30% khi chi phí sửa lớn / Đặt linh kiện',
        'Báo giá sửa chữa > 1.000.000 đ hoặc linh kiện đặc biệt',
        '1. Tính tổng chi phí dự kiến: 2.000.000 đ\n2. Thu ngân thu tiền cọc 30%\n3. Ghi nhận vào phiếu sửa',
        'Tổng chi phí: 2.000.000 đ -> Tiền cọc 30%: 600.000 đ',
        'Lưu tienDatCoc = 600.000 đ, daDatCoc = True. Phiếu chuyển sang ChoLinhKien hoặc DangSua.',
        'Critical', 'PASS', 'Đã xác minh schema MySQL & logic cọc'
    ],
    [
        'TC_STD_03', 'Khách Duyệt Giá', 'Khách hàng duyệt báo giá trực tuyến (Online Approval)',
        'Phiếu ở trạng thái DangKiemTra đã có bóc tách giá',
        '1. Khách mở link tra cứu trực tuyến\n2. Xem bảng bóc tách chi phí\n3. Bấm nút [Đồng Ý Báo Giá]',
        'Mã phiếu: PS01, Khách bấm Đồng Ý',
        'Hệ thống ghi nhận duyetBaoGia = "DONG_Y", KTV nhận thông báo được phép tiến hành thay thế.',
        'High', 'PASS', 'Đã test renderQuoteApprovalBox'
    ],
    [
        'TC_STD_04', 'Kỹ Thuật Viên', 'KTV nhận việc và chuyển trạng thái sửa chữa',
        'Phiếu ở trạng thái KhoiTao hoặc DangKiemTra',
        '1. KTV đăng nhập hệ thống\n2. Vào tab "Máy Chờ Nhận"\n3. Bấm [Nhận Xử Lý]',
        'KTV: Lê Văn Cường (NV03), Phiếu: PS03',
        'Gán maKTV = NV03, chuyển trạng thái sang DangKiemTra / DangSua, hiển thị trong danh sách máy của tôi.',
        'High', 'PASS', 'Endpoint /phieu-sua/{ma}/nhan-viec OK'
    ],
    [
        'TC_STD_05', 'Kho Linh Kiện', 'Gán linh kiện vào phiếu & Tự động trừ tồn kho',
        'Linh kiện có sẵn trong kho (soLuongTon > 0)',
        '1. KTV mở modal xử lý phiếu\n2. Chọn linh kiện LK01 (Pin) số lượng 1\n3. Bấm [Thêm Linh Kiện]',
        'Mã LK: LK01, Số lượng: 1, Đơn giá: 450.000 đ',
        'Ghi nhận vào chi_tiet_phieu_sua_chua, trừ tồn kho LK01 đi 1, tổng tiền linh kiện tự động cộng 450.000 đ.',
        'Critical', 'PASS', 'Tự động kiểm tra FK integrity kho LK'
    ],
    [
        'TC_STD_06', 'Thông Báo Tự Động', 'Báo máy đã sửa xong & Sinh tin nhắn Zalo/SMS',
        'KTV hoàn tất sửa chữa và kiểm tra chức năng đạt chuẩn',
        '1. Chuyển trạng thái sang DaSuaXong\n2. Hệ thống gọi AI soạn tin nhắn theo kênh Zalo/SMS\n3. Gửi thông báo đến khách',
        'Phiếu: PS01, Kênh: Zalo OA',
        'Sinh tin nhắn chúc mừng thân thiện kèm mã tra cứu và link bảo hành điện tử, sẵn sàng gửi khách.',
        'Medium', 'PASS', 'Hàm generate_customer_message OK'
    ],
    [
        'TC_STD_07', 'Thu Ngân 70%', 'Thanh toán 70% còn lại sau khi trừ tiền cọc 30%',
        'Thiết bị đã sửa xong (DaSuaXong), khách đến nhận máy',
        '1. Thu ngân mở phiếu thanh toán\n2. Hệ thống hiển thị: Tổng 2.000.000 đ - Đã cọc 600.000 đ\n3. Thu nốt 1.400.000 đ (70%)',
        'Số tiền thực thu: 1.400.000 đ. Phương thức: Chuyển khoản VietQR',
        'Cập nhật daThanhToan = True, trangThai = DaThanhToan. Xuất hóa đơn VAT & biên lai tất toán.',
        'Critical', 'PASS', 'Đã kiểm thử logic trừ cọc chính xác'
    ],
    [
        'TC_STD_08', 'Bảo Hành Điện Tử', 'Tự động kích hoạt E-Warranty 6T/12T chuẩn hãng',
        'Phiếu đã thanh toán và bàn giao thành công',
        '1. Hệ thống quét các linh kiện đã thay\n2. Tự động sinh Thẻ bảo hành điện tử cho từng linh kiện\n3. Tính ngày hết hạn',
        'Pin: BH 12 tháng (1 đổi 1)\nMàn hình: BH 6 tháng',
        'Kích hoạt e-warranty tra cứu trực tuyến 24/7 theo số điện thoại khách hàng.',
        'High', 'PASS', 'Module bảo hành chi tiết OK'
    ],
    [
        'TC_STD_09', 'Đánh Giá KTV', 'Khách hàng chấm điểm 5 sao & Đánh giá dịch vụ',
        'Khách đã nhận máy và thanh toán hoàn tất',
        '1. Khách mở cổng tra cứu khach_hang.html\n2. Chọn 5 sao cho KTV phụ trách\n3. Viết cảm nhận và bấm [Gửi Đánh Giá]',
        'Số sao: 5/5, Nhận xét: "KTV tư vấn tận tâm, thay pin zin dùng rất trâu!"',
        'Lưu vào bảng danh_gia, tính điểm KPI trung bình cho Kỹ thuật viên phụ trách.',
        'Medium', 'PASS', 'API /api/khach-hang/danh-gia OK'
    ],
    [
        'TC_STD_10', 'Khiếu Nại Kỹ Thuật', 'Gửi khiếu nại dịch vụ & Tiếp nhận xử lý sự cố',
        'Khách phát hiện lỗi tái phát hoặc chưa hài lòng',
        '1. Bấm [Gửi Khiếu Nại / Góp Ý]\n2. Chọn loại: "Khiếu nại kỹ thuật"\n3. Nhập nội dung và gửi',
        'Nội dung: "Màn hình bị hở mép viền trên sau 2 ngày sử dụng"',
        'Tạo bản ghi khieu_nai, thông báo ngay cho Quản lý & KTV trưởng để liên hệ hỗ trợ bảo hành ưu tiên trong 2h.',
        'High', 'PASS', 'Modal khiếu nại hoạt động hoàn hảo'
    ]
]
format_data_rows(ws2, tc_core)
file2_path = os.path.abspath('docs/test_cases/Bang_Test_Case_Quy_Chuan_Sua_Chua_Va_Dat_Coc.xlsx')
wb2.save(file2_path)
print(f"Generated: {file2_path}")
print("ALL DONE SUCCESSFULLY!")
